#!/usr/bin/env python3

from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path.home() / "postexport-kinetics" / "real_datasets"

INPUT = BASE / "final_real_data_audit" / "final_all_realdata_standardized.tsv"
OUTPUT = BASE / "S5testresults_final.xlsx"

df = pd.read_csv(INPUT, sep="\t")

required = [
    "dataset_final",
    "gene_final",
    "event_final",
    "p_final",
    "q_final",
    "sigma_c_final",
    "IR_final",
]

missing = [x for x in required if x not in df.columns]
if missing:
    raise RuntimeError(f"Missing required columns: {missing}")

dataset_order = ["Kc167", "K562", "NIH-3T3", "mESC"]

# ---------------------------------------------------------------------
# Derived fields used only for reporting
# ---------------------------------------------------------------------

df["half_time_min"] = pd.NA

mask = df["sigma_c_final"].notna() & (df["sigma_c_final"] > 0)
df.loc[mask, "half_time_min"] = (
    0.6931471805599453 / df.loc[mask, "sigma_c_final"]
)

df["nominal_p_lt_005"] = df["p_final"] < 0.05
df["BH_q_lt_010"] = df["q_final"] < 0.10
df["BH_q_lt_005"] = df["q_final"] < 0.05
df["sigma_c_boundary"] = (
    df["sigma_c_final"].notna()
    & (df["sigma_c_final"].abs() < 1e-12)
)

# ---------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------

summary_rows = []

for ds in dataset_order:
    x = df[df["dataset_final"] == ds].copy()

    summary_rows.append({
        "Dataset": ds,
        "Tested RI events": len(x),
        "Unique genes": x["gene_final"].dropna().nunique(),
        "p < 0.05": int((x["p_final"] < 0.05).sum()),
        "q < 0.10": int((x["q_final"] < 0.10).sum()),
        "q < 0.05": int((x["q_final"] < 0.05).sum()),
        "sigma_c boundary": int(x["sigma_c_boundary"].sum()),
        "Boundary fraction": float(x["sigma_c_boundary"].mean()),
    })

summary = pd.DataFrame(summary_rows)

# ---------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)

sub_fill = PatternFill("solid", fgColor="D9EAF7")

def write_dataframe(ws, data):
    # Header
    for col_idx, col_name in enumerate(data.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Rows
    for row_idx, row in enumerate(data.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            ws.cell(row=row_idx, column=col_idx, value=value)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Column widths
    for idx, col_name in enumerate(data.columns, start=1):
        max_len = len(str(col_name))

        for cell in ws[get_column_letter(idx)][1:]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))

        width = min(max(max_len + 2, 10), 55)
        ws.column_dimensions[get_column_letter(idx)].width = width


# ---------------------------------------------------------------------
# Summary sheet
# ---------------------------------------------------------------------

ws = wb.create_sheet("Summary")

for col_idx, col_name in enumerate(summary.columns, start=1):
    c = ws.cell(row=1, column=col_idx, value=col_name)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center")

for row_idx, row in enumerate(summary.itertuples(index=False), start=2):
    for col_idx, value in enumerate(row, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)

# Number formats
boundary_col = summary.columns.get_loc("Boundary fraction") + 1
for r in range(2, ws.max_row + 1):
    ws.cell(r, boundary_col).number_format = "0.0%"

ws.freeze_panes = "A2"

for i, col in enumerate(summary.columns, start=1):
    ws.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 3)


# ---------------------------------------------------------------------
# One sheet per dataset
# ---------------------------------------------------------------------

export_cols = [
    "dataset_final",
    "gene_final",
    "event_final",
    "p_final",
    "q_final",
    "sigma_c_final",
    "half_time_min",
    "IR_final",
    "nominal_p_lt_005",
    "BH_q_lt_010",
    "BH_q_lt_005",
    "sigma_c_boundary",
]

for ds in dataset_order:
    x = df[df["dataset_final"] == ds][export_cols].copy()

    # Sort most statistically supported first
    x = x.sort_values(
        ["q_final", "p_final", "IR_final"],
        ascending=[True, True, False],
        na_position="last",
    )

    ws = wb.create_sheet(ds)
    write_dataframe(ws, x)

    # Numeric formatting
    col_map = {name: i + 1 for i, name in enumerate(x.columns)}

    for r in range(2, ws.max_row + 1):
        ws.cell(r, col_map["p_final"]).number_format = "0.000000"
        ws.cell(r, col_map["q_final"]).number_format = "0.000000"
        ws.cell(r, col_map["sigma_c_final"]).number_format = "0.000000"
        ws.cell(r, col_map["half_time_min"]).number_format = "0.0"
        ws.cell(r, col_map["IR_final"]).number_format = "0.000"


# ---------------------------------------------------------------------
# Dedicated mESC FDR-supported sheet
# ---------------------------------------------------------------------

mesc = df[
    (df["dataset_final"] == "mESC")
    & (df["q_final"] < 0.05)
][export_cols].copy()

mesc = mesc.sort_values(
    ["q_final", "p_final", "IR_final"],
    ascending=[True, True, False],
)

ws = wb.create_sheet("mESC_FDR05")
write_dataframe(ws, mesc)

for col_idx, col_name in enumerate(mesc.columns, start=1):
    ws.cell(1, col_idx).fill = sub_fill
    ws.cell(1, col_idx).font = Font(bold=True)

# ---------------------------------------------------------------------
# Metadata sheet
# ---------------------------------------------------------------------

ws = wb.create_sheet("README")

readme = [
    ["File", "S5testresults_final.xlsx"],
    [
        "Description",
        "Complete final event-level constrained nested-model results across real datasets."
    ],
    [
        "Source",
        "final_real_data_audit/final_all_realdata_standardized.tsv"
    ],
    [
        "Multiple testing",
        "Benjamini-Hochberg correction performed separately within each dataset."
    ],
    [
        "p_final",
        "Final bootstrap p-value."
    ],
    [
        "q_final",
        "Dataset-specific Benjamini-Hochberg adjusted q-value."
    ],
    [
        "sigma_c_final",
        "Estimated post-export conversion rate."
    ],
    [
        "half_time_min",
        "log(2) / sigma_c_final when sigma_c_final > 0."
    ],
    [
        "IR_final",
        "Relative RSS improvement of the full model over the constrained null."
    ],
    [
        "sigma_c_boundary",
        "TRUE when the fitted post-export conversion rate is at the non-negativity boundary."
    ],
]

for r, row in enumerate(readme, start=1):
    ws.cell(r, 1, row[0])
    ws.cell(r, 2, row[1])

ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 95
ws["A1"].font = Font(bold=True)
ws["B1"].font = Font(bold=True)

# ---------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------

wb.save(OUTPUT)

print("\n============================================================")
print("FINAL SUPPLEMENTARY TABLE S5 GENERATED")
print("============================================================")
print(f"Output: {OUTPUT}")
print("\nSummary:")
print(summary.to_string(index=False))